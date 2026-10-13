import unittest
from datetime import date, datetime, timezone
from unittest.mock import patch
from io import BytesIO
from uuid import uuid4
from openpyxl import load_workbook
from sqlalchemy import select
import test_billing as fixture
from app.models import LedgerTransaction
from app.services.transaction_export_service import totals, pdf as render_pdf


class ExportTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = fixture.BillingTests.asyncSetUp
    asyncTearDown = fixture.BillingTests.asyncTearDown
    auth_headers = fixture.BillingTests.auth_headers

    async def seed(self):
        rows = []
        async with self.sessions() as db:
            for index, (kind, category, amount) in enumerate([
                ('income', 'salary', 1000000), ('expense', 'food', 25000), ('income', 'cash_withdrawal', 500000),
                ('income', 'receivable', 100000), ('transfer', 'savings', 300000),
            ]):
                row = LedgerTransaction(id=uuid4(), user_id=self.user_id, kind=kind, category=category,
                    description='=HYPERLINK("https://example.test")' if index == 1 else 'Synthetic ' + str(index),
                    amount=amount, transaction_date=date(2026, 10, index + 1), notes='@SUM(1,2)',
                    source_label='BLU · ********6201', destination_label='')
                db.add(row); rows.append(row)
            foreign = LedgerTransaction(id=uuid4(), user_id=self.other_id, kind='expense', category='other',
                description='Private other user', amount=999, transaction_date=date(2026, 10, 1), notes='', source_label='', destination_label='')
            db.add(foreign)
            await db.commit()
        return rows, foreign

    def payload(self, rows, **changes):
        return {'Month': '2026-10', 'TransactionIds': [str(row.id) for row in rows]} | changes

    async def test_export_order_uses_transaction_date_then_created_clock_in_wib(self):
        rows, _ = await self.seed()
        recorded = [datetime(2026, 10, 13, 20, tzinfo=timezone.utc),
                    datetime(2026, 10, 12, 3, tzinfo=timezone.utc),
                    datetime(2026, 10, 11, 15, tzinfo=timezone.utc),
                    datetime(2026, 10, 14, 3, tzinfo=timezone.utc),
                    datetime(2026, 10, 10, 0, tzinfo=timezone.utc)]
        async with self.sessions() as db:
            for index, row in enumerate(rows):
                stored = await db.get(LedgerTransaction, row.id)
                stored.transaction_date = date(2026, 10, 2 if index == 4 else 1)
                stored.created_at = recorded[index]
            await db.commit()
        selected = [rows[index] for index in (2, 4, 3, 1, 0)]
        expected = [rows[index] for index in (0, 3, 1, 2, 4)]
        payload = self.payload(selected)
        result = await self.client.post('/transactions/export/xlsx', headers=self.headers, json=payload)
        self.assertEqual(result.status_code, 200)
        sheet = load_workbook(BytesIO(result.content)).active
        self.assertEqual([sheet.cell(index, 2).value for index in range(13, 18)], [row.description for row in expected])
        with patch('app.services.transaction_export_service.pdf', wraps=render_pdf) as rendered:
            result = await self.client.post('/transactions/export/pdf', headers=self.headers, json=payload)
            self.assertEqual(result.status_code, 200)
            self.assertEqual([row.id for row in rendered.call_args.args[0]], [row.id for row in expected])

    async def test_xlsx_numeric_dates_formulas_cache_and_no_formula_injection(self):
        rows, _ = await self.seed()
        result = await self.client.post('/transactions/export/xlsx', headers=self.headers, json=self.payload(list(reversed(rows))))
        self.assertEqual(result.status_code, 200, result.text if result.status_code != 200 else '')
        self.assertEqual(result.headers['cache-control'], 'no-store')
        wb = load_workbook(BytesIO(result.content))
        ws = wb.active
        self.assertEqual([ws.cell(index, 1).value.date() for index in range(13, 18)], [row.transaction_date for row in rows])
        self.assertEqual(ws['E13'].value, 1000000)
        self.assertEqual(ws['E14'].value, -25000)
        self.assertEqual(ws['E13'].font.color.rgb, '0015803D')
        self.assertEqual(ws['D14'].font.color.rgb, '00B91C1C')
        self.assertEqual(ws['E14'].font.color.rgb, '00B91C1C')
        self.assertEqual(ws['E13'].data_type, 'n')
        self.assertEqual(ws['B14'].value, '=HYPERLINK("https://example.test")')
        self.assertEqual(ws['B14'].data_type, 's')
        self.assertEqual(ws['H13'].data_type, 's')
        self.assertEqual(ws.freeze_panes, 'C13')
        self.assertIn('Transactions', ws.tables)
        self.assertEqual(ws['B6'].data_type, 'f')
        self.assertEqual(ws['H6'].value, '=B6+B7+E6')
        cached = load_workbook(BytesIO(result.content), data_only=True).active
        expected = totals(rows)
        expected[1] = -expected[1]
        self.assertEqual([cached[cell].value for cell in ('B6', 'E6', 'H6', 'B7', 'E7', 'H7')], expected)
        self.assertEqual(cached['H6'].value, 1075000)
        self.assertEqual(len(list(ws.rows)), 17)

    async def test_pdf_filters_and_read_only_owner_month_void_validation(self):
        rows, foreign = await self.seed()
        payload = self.payload([rows[1]], Search='makan', Kind='expense', Bank='BLU · ********6201')
        pdf = await self.client.post('/transactions/export/pdf', headers=self.headers, json=payload)
        self.assertEqual(pdf.status_code, 200, pdf.text if pdf.status_code != 200 else '')
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        self.assertEqual(pdf.headers['content-type'], 'application/pdf')
        for invalid in (self.payload([foreign]), self.payload(rows, Month='2026-11'), self.payload(rows, Kind='income')):
            self.assertEqual((await self.client.post('/transactions/export/pdf', headers=self.headers, json=invalid)).status_code, 409)
        self.assertEqual((await self.client.post('/transactions/export/pdf', json=self.payload(rows))).status_code, 401)
        async with self.sessions() as db:
            item = await db.get(LedgerTransaction, rows[1].id)
            from datetime import datetime, timezone
            item.voided_at = datetime.now(timezone.utc)
            await db.commit()
        self.assertEqual((await self.client.post('/transactions/export/pdf', headers=self.headers, json=payload)).status_code, 409)
        for format in ('pdf', 'xlsx'):
            self.assertEqual((await self.client.post('/transactions/export/' + format, headers=self.headers, json=self.payload([]))).status_code, 422)
            duplicated = self.payload([rows[0], rows[0]])
            self.assertEqual((await self.client.post('/transactions/export/' + format, headers=self.headers, json=duplicated)).status_code, 422)
        self.assertEqual((await self.client.post('/transactions/export/csv', headers=self.headers, json=self.payload(rows))).status_code, 422)
        async with self.sessions() as db:
            persisted = list(await db.scalars(select(LedgerTransaction)))
            self.assertEqual(len(persisted), 6)
            self.assertEqual(sum(item.amount for item in persisted), sum(item.amount for item in rows) + foreign.amount)
