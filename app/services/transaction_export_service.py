"""Read-only reports from authenticated, active ledger rows selected by the UI."""
from calendar import monthrange
from datetime import date, timedelta, timezone
from io import BytesIO
from pathlib import Path
import re
from xml.sax.saxutils import escape
from fastapi import HTTPException
from starlette.concurrency import run_in_threadpool
from sqlalchemy import select
from app.models.routine_model import LedgerTransaction
from app.schemas.transaction_export_schema import TransactionExportInput

CATEGORIES = {
    'salary': 'Gaji & pendapatan', 'food': 'Makan & minum', 'transport': 'Transportasi',
    'shopping': 'Belanja', 'bills': 'Tagihan', 'health': 'Kesehatan', 'other': 'Lainnya',
    'family': 'Keluarga / Setoran', 'savings': 'Alokasi Tabungan', 'top_up': 'Top Up',
    'entertainment': 'Hiburan / Gaya Hidup', 'education': 'Pendidikan / Pengembangan Diri',
    'donation': 'Sedekah / Donasi / Hadiah', 'debt': 'Cicilan / Hutang', 'bonus': 'Bonus',
    'investment': 'Hasil Investasi', 'sales': 'Hasil Penjualan', 'receivable': 'Pengembalian Piutang',
    'cash_withdrawal': 'Tarik Tunai',
}
KINDS = {'income': 'Pemasukan', 'expense': 'Pengeluaran', 'transfer': 'Transfer Antar Rekening'}
HEADERS = ['Tanggal', 'Deskripsi', 'Kategori', 'Jenis Transaksi', 'Nominal (Rp)', 'Rekening Sumber', 'Rekening Tujuan', 'Catatan']
NOTE = 'Tarik tunai dan transfer rekening sendiri bukan pendapatan. Arus kas bersih = pendapatan + pengembalian piutang - pengeluaran.'


def safe_text(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))


def money(value):
    return ('-' if value < 0 else '') + 'Rp ' + f'{abs(value):,}'.replace(',', '.')


def totals(rows):
    income = sum(int(row.amount) for row in rows if row.kind == 'income' and row.category not in ('receivable', 'cash_withdrawal'))
    received = sum(int(row.amount) for row in rows if row.kind == 'income' and row.category == 'receivable')
    expense = sum(int(row.amount) for row in rows if row.kind == 'expense')
    return [income, expense, income + received - expense, received,
        sum(int(row.amount) for row in rows if row.kind == 'transfer'),
        sum(int(row.amount) for row in rows if row.kind == 'income' and row.category == 'cash_withdrawal')]


def filters(payload):
    return 'Pencarian: ' + (payload.search or 'Semua') + ' | Jenis: ' + KINDS.get(payload.kind, 'Semua') + ' | Rekening: ' + (payload.bank or 'Semua')


async def load(payload: TransactionExportInput, user_id, db):
    start = date.fromisoformat(payload.month + '-01')
    end = start.replace(day=monthrange(start.year, start.month)[1])
    rows = list(await db.scalars(select(LedgerTransaction).where(
        LedgerTransaction.user_id == user_id, LedgerTransaction.voided_at.is_(None),
        LedgerTransaction.id.in_(payload.transaction_ids), LedgerTransaction.transaction_date.between(start, end))))
    if len(rows) != len(payload.transaction_ids):
        raise HTTPException(409, 'Data transaksi berubah atau tidak tersedia. Muat ulang sebelum mengekspor.')
    # Re-check filters so the context printed in the report matches its contents.
    for row in rows:
        category = CATEGORIES.get(row.category, row.category)
        if ((payload.kind != 'all' and row.kind != payload.kind)
            or (payload.bank and payload.bank not in (row.source_label, row.destination_label))
            or payload.search.lower() not in (row.description + ' ' + category).lower()):
            raise HTTPException(409, 'Filter transaksi berubah. Muat ulang sebelum mengekspor.')
    positions = {key: index for index, key in enumerate(payload.transaction_ids)}
    def order(row: LedgerTransaction):
        created = row.created_at
        if created.tzinfo is None:
            created = created.replace(tzinfo=timezone.utc)
        # Use the recorded clock time in WIB; the transaction date remains authoritative.
        clock = created.astimezone(timezone(timedelta(hours=7))).time()
        return row.transaction_date.toordinal(), clock, positions[row.id]
    return sorted(rows, key=order)


def excel(rows, payload, name):
    # Application exports run in Python/Vercel using the existing openpyxl dependency.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.worksheet.worksheet import Worksheet
    from openpyxl.worksheet.properties import PageSetupProperties
    from openpyxl.workbook.properties import CalcProperties
    wb = Workbook()
    ws = wb.active
    if not isinstance(ws, Worksheet):
        raise RuntimeError('Worksheet untuk ekspor transaksi tidak tersedia.')
    ws.title = 'Transaksi'
    def text(cell, value):
        cell.value = safe_text(value)
        cell.data_type = 's'  # User strings must never become Excel formulas.
    for row, value in [(1, 'Laporan Transaksi'), (2, f'Periode {payload.month} | {name}'),
                       (3, f'{len(rows)} transaksi sesuai tabel yang diekspor'), (4, filters(payload)),
                       (9, NOTE), (10, 'Perubahan pada file ini tidak mengubah data di aplikasi.')]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        text(ws.cell(row, 1), value)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical='center')
        ws.cell(row, 1).font = Font(name='Arial', size=11 if row != 1 else 16, bold=row == 1, color='334155')
        ws.row_dimensions[row].height = 32 if row in (1, 4, 9) else 24
    end = 12 + len(rows)
    amount, kind, category = f'E13:E{end}', f'D13:D{end}', f'C13:C{end}'
    formulas = [
        ('A6', 'Pendapatan', 'B6', f'=SUMIFS({amount},{kind},"Pemasukan",{category},"<>Tarik Tunai",{category},"<>Pengembalian Piutang")'),
        ('D6', 'Pengeluaran', 'E6', f'=SUMIF({kind},"Pengeluaran",{amount})'),
        ('G6', 'Arus Kas Bersih', 'H6', '=B6+B7+E6'),
        ('A7', 'Pengembalian Piutang', 'B7', f'=SUMIFS({amount},{kind},"Pemasukan",{category},"Pengembalian Piutang")'),
        ('D7', 'Transfer Antar Rekening', 'E7', f'=SUMIF({kind},"Transfer Antar Rekening",{amount})'),
        ('G7', 'Tarik Tunai', 'H7', f'=SUMIFS({amount},{kind},"Pemasukan",{category},"Tarik Tunai")'),
    ]
    for label_cell, label, value_cell, formula in formulas:
        text(ws[label_cell], label)
        color = '15803D' if value_cell in ('B6', 'B7') else 'B91C1C' if value_cell == 'E6' else '334155'
        ws[label_cell].font = Font(name='Arial', bold=True, color=color)
        ws[label_cell].alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[ws[label_cell].row].height = 32
        ws[value_cell] = formula
        ws[value_cell].number_format = '"Rp "#,##0;[Red]-"Rp "#,##0'
        ws[value_cell].font = Font(name='Arial', bold=True, color=color)
        ws[value_cell].alignment = Alignment(horizontal='right', vertical='center', indent=1)
    for operator, color in [('greaterThan', '15803D'), ('lessThan', 'B91C1C')]:
        ws.conditional_formatting.add('H6', CellIsRule(operator=operator, formula=['0'], font=Font(color=color)))
    for column, header in enumerate(HEADERS, 1):
        text(ws.cell(12, column), header)
        ws.cell(12, column).font = Font(name='Arial', bold=True, color='FFFFFF')
        ws.cell(12, column).fill = PatternFill('solid', fgColor='334155')
    widths = [14, 44, 30, 26, 28, 36, 36, 80]
    from openpyxl.utils import get_column_letter
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    for index, row in enumerate(rows, 13):
        values = [row.transaction_date, row.description, CATEGORIES.get(row.category, row.category), KINDS[row.kind],
                  -abs(int(row.amount)) if row.kind == 'expense' else int(row.amount), row.source_label or 'Tunai / Tidak dicatat', row.destination_label or 'Tunai / Tidak dicatat', row.notes]
        for column, value in enumerate(values, 1):
            cell = ws.cell(index, column)
            if column in (1, 5):
                cell.value = value
            else:
                text(cell, value)
            cell.alignment = Alignment(horizontal='right' if column == 5 else 'left', vertical='top', indent=1, wrap_text=column not in (1, 5))
            color = ('15803D' if row.kind == 'income' else 'B91C1C' if row.kind == 'expense' else '334155') if column in (4, 5) else '334155'
            cell.font = Font(name='Arial', size=11, bold=column in (4, 5), color=color)
        ws.cell(index, 1).number_format = 'dd-mm-yyyy'
        ws.cell(index, 5).number_format = '"Rp "#,##0;[Red]-"Rp "#,##0'
        ws.row_dimensions[index].height = min(409, 30 + 14 * max(len(row.description)//38, len(row.notes or '')//68,
            len(row.source_label or '')//30, len(row.destination_label or '')//30,
            (row.description or '').count('\n'), (row.notes or '').count('\n')))
    table = Table(displayName='Transactions', ref=f'A12:H{end}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = 'C13'
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = '12:12'
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_area = f'A1:H{end}'
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
    output = BytesIO()
    wb.save(output)
    # Cached values keep the initial summary readable in non-calculating viewers;
    # Excel retains the formulas and recalculates them after edits.
    from zipfile import ZipFile, ZIP_DEFLATED
    from xml.etree import ElementTree as ET
    summary_values = totals(rows)
    summary_values[1] = -abs(summary_values[1])
    cached = dict(zip(['B6', 'E6', 'H6', 'B7', 'E7', 'H7'], summary_values))
    result = BytesIO()
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with ZipFile(output) as source, ZipFile(result, 'w', ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                root = ET.fromstring(content)
                for cell in root.iter(ns + 'c'):
                    reference = cell.get('r')
                    if reference is not None and reference in cached:
                        value = cell.find(ns + 'v')
                        if value is not None:
                            value.text = str(cached[reference])
                content = ET.tostring(root, encoding='utf-8')
            target.writestr(info, content)
    return result.getvalue()


def pdf(rows, payload, name):
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable
    fonts = Path(reportlab.__file__).parent / 'fonts'
    for font, filename in [('TransactionRegular', 'Vera.ttf'), ('TransactionBold', 'VeraBd.ttf')]:
        if font not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font, str(fonts / filename)))
    regular = ParagraphStyle('Transaction', fontName='TransactionRegular', fontSize=8, leading=11, textColor=colors.HexColor('#334155'))
    bold = ParagraphStyle('TransactionBold', parent=regular, fontName='TransactionBold')
    income_style = ParagraphStyle('TransactionIncome', parent=bold, textColor=colors.HexColor('#15803D'))
    expense_style = ParagraphStyle('TransactionExpense', parent=bold, textColor=colors.HexColor('#B91C1C'))
    title = ParagraphStyle('TransactionTitle', parent=bold, fontSize=20, leading=24)
    def p(value, style=regular):
        return Paragraph(escape(safe_text(value)).replace('\n', '<br/>'), style)
    story = [p('Laporan Transaksi', title), Spacer(1, 8), p(f'Periode {payload.month} | {name}'),
        p(f'{len(rows)} transaksi sesuai tabel yang diekspor'), p(filters(payload)), Spacer(1, 12)]
    labels = ['Pendapatan', 'Pengeluaran', 'Arus Kas Bersih', 'Pengembalian Piutang', 'Transfer Antar Rekening', 'Tarik Tunai']
    values = totals(rows)
    summary_styles = [income_style, expense_style,
        income_style if values[2] > 0 else expense_style if values[2] < 0 else bold,
        income_style, bold, bold]
    summary_values = values.copy()
    summary_values[1] = -abs(summary_values[1])
    summary = LongTable([[p(labels[i] + '\n' + money(summary_values[i]), summary_styles[i]) for i in range(3)],
                         [p(labels[i] + '\n' + money(summary_values[i]), summary_styles[i]) for i in range(3, 6)]], colWidths=[(A4[0]-64)/3]*3)
    summary.setStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#EAF0FF')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 9), ('BOTTOMPADDING', (0, 0), (-1, -1), 9)])
    story.extend([summary, Spacer(1, 8), p(NOTE), Spacer(1, 12)])
    headers = ['Tanggal', 'Deskripsi / Catatan', 'Kategori / Jenis', 'Nominal', 'Rekening']
    header_style = ParagraphStyle('TransactionHeader', parent=bold, textColor=colors.white)
    data: list[list[Paragraph | list[Paragraph]]] = [[p(header, header_style) for header in headers]]
    for row in rows:
        description = row.description + ('\nCatatan: ' + row.notes if row.notes else '')
        amount = int(row.amount)
        transaction_style = income_style if row.kind == 'income' else expense_style if row.kind == 'expense' else bold
        displayed_amount = -abs(amount) if row.kind == 'expense' else amount
        accounts = 'Sumber: ' + (row.source_label or 'Tunai / Tidak dicatat') + '\nTujuan: ' + (row.destination_label or 'Tunai / Tidak dicatat')
        data.append([p(row.transaction_date.strftime('%d-%m-%Y')), p(description),
            [p(CATEGORIES.get(row.category, row.category)), p(KINDS[row.kind], transaction_style)],
            p(money(displayed_amount), transaction_style), p(accounts)])
    table = LongTable(data, colWidths=[62, 145, 96, 100, A4[0]-64-403], repeatRows=1, splitInRow=1)
    table.setStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7F9FC')]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)])
    story.append(table)
    output = BytesIO()
    width, height = A4
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('TransactionRegular', 8)
        canvas.setFillColor(colors.HexColor('#64748B'))
        canvas.drawString(32, 18, 'Finance Tracker | Laporan Transaksi')
        canvas.drawRightString(width-32, 18, f'Halaman {doc.page}')
        canvas.restoreState()
    SimpleDocTemplate(output, pagesize=(width, height), leftMargin=32, rightMargin=32, topMargin=28, bottomMargin=38,
        title='Laporan Transaksi', author='Finance Tracker').build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


async def export(payload, format, user, db):
    rows = await load(payload, user.id, db)
    if format == 'pdf':
        return await run_in_threadpool(pdf, rows, payload, user.name), 'application/pdf'
    return await run_in_threadpool(excel, rows, payload, user.name), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
